from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = r"C:\Users\dahli\Documents\机器人算法开发\麦克纳姆底盘横移偏移排查手册.docx"
BLUE = "2E74B5"
DARK = "1F4D78"
LIGHT = "E8EEF5"
PALE = "F4F6F9"
GOLD = "FFF3CD"
RED = "FCE8E6"
GRAY = "666666"

def set_font(run, size=11, bold=False, color="000000", name="Microsoft YaHei"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def margins(cell, top=90, start=120, bottom=90, end=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn("w:" + tag))
        if node is None:
            node = OxmlElement("w:" + tag)
            tcMar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")

def set_table_widths(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths))); tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd"); tblPr.append(tblInd)
    tblInd.set(qn("w:w"), "120"); tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for w in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(w)); grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tcW = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW"); cell._tc.get_or_add_tcPr().append(tcW)
            tcW.set(qn("w:w"), str(widths[i])); tcW.set(qn("w:type"), "dxa")
            margins(cell)

def format_cell(cell, bold=False, color="000000", align=WD_ALIGN_PARAGRAPH.LEFT, size=9.5):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        for r in p.runs: set_font(r, size, bold, color)

def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h; shade(t.rows[0].cells[i], LIGHT)
        format_cell(t.rows[0].cells[i], True, DARK, WD_ALIGN_PARAGRAPH.CENTER)
    trPr = t.rows[0]._tr.get_or_add_trPr(); rep = OxmlElement("w:tblHeader"); rep.set(qn("w:val"), "true"); trPr.append(rep)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value); format_cell(cells[i], False, align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
    set_table_widths(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def p(doc, text="", bold_prefix=None, after=6, style=None):
    para = doc.add_paragraph(style=style)
    para.paragraph_format.space_after = Pt(after)
    para.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        a = para.add_run(bold_prefix); set_font(a, 11, True, DARK)
        b = para.add_run(text[len(bold_prefix):]); set_font(b)
    else:
        r = para.add_run(text); set_font(r)
    return para

def bullet(doc, text):
    para = doc.add_paragraph(style="List Bullet")
    para.paragraph_format.left_indent = Inches(0.375)
    para.paragraph_format.first_line_indent = Inches(-0.188)
    para.paragraph_format.space_after = Pt(4); para.paragraph_format.line_spacing = 1.25
    set_font(para.add_run(text))

def step(doc, title, body):
    para = doc.add_paragraph(style="List Number")
    para.paragraph_format.left_indent = Inches(0.375)
    para.paragraph_format.first_line_indent = Inches(-0.188)
    para.paragraph_format.space_after = Pt(4); para.paragraph_format.line_spacing = 1.25
    set_font(para.add_run(title + "："), 11, True, DARK)
    set_font(para.add_run(body))

def heading(doc, text, level=1):
    para = doc.add_paragraph(text, style=f"Heading {level}")
    para.paragraph_format.keep_with_next = True
    return para

def callout(doc, label, text, fill=PALE):
    t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.LEFT
    c = t.cell(0,0); shade(c, fill); margins(c, 130, 160, 130, 160)
    pp = c.paragraphs[0]; pp.paragraph_format.space_after = Pt(0); pp.paragraph_format.line_spacing = 1.2
    set_font(pp.add_run(label + "  "), 10.5, True, DARK)
    set_font(pp.add_run(text), 10.5)
    set_table_widths(t, [9360]); doc.add_paragraph().paragraph_format.space_after = Pt(2)

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]; normal.font.name = "Microsoft YaHei"; normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(11); normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.25
for name, size, color, before, after in [("Heading 1",16,BLUE,18,10),("Heading 2",13,BLUE,14,7),("Heading 3",12,DARK,10,5)]:
    s=styles[name]; s.font.name="Microsoft YaHei"; s._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei")
    s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color)
    s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True

header = sec.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(header.add_run("底盘调试参考｜Four_Motor_PID_Test"), 8.5, False, GRAY)
footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); footer._p.append(fld)

title = doc.add_paragraph(); title.paragraph_format.space_before=Pt(30); title.paragraph_format.space_after=Pt(5)
set_font(title.add_run("麦克纳姆底盘横移偏移排查手册"), 24, True, DARK)
sub = doc.add_paragraph(); sub.paragraph_format.space_after=Pt(18)
set_font(sub.add_run("适用于 STM32 四电机速度闭环工程 Four_Motor_PID_Test"), 12, False, GRAY)
callout(doc, "核心判断", "悬空轮速接近、落地偏移明显，首先怀疑负载下轮速跟随不足或麦轮机械受力差异，不要先盲目增大 PID。", GOLD)

heading(doc, "1  先理解为什么悬空正常、落地却偏", 1)
p(doc, "悬空时，电机只需克服轴承和减速箱阻力，较小 PWM 就能达到目标转速。落地横移时，麦轮滚子、地面摩擦和车体重量都会增加负载。如果某个轮子的负载更大，它的实际转速就会低于目标转速，四轮合力失衡后，底盘便会斜走或转头。")
callout(doc, "排查目标", "找到落地后哪一个轮子的 actual_rpm 跟不上 target_rpm，并判断是控制输出不足，还是机械阻力/打滑造成。")

heading(doc, "2  正常横移时四轮应该怎样转", 1)
p(doc, "工程中的轮号定义为 M1 左前、M2 右前、M3 左后、M4 右后。遥控器向右横移时，目标轮速符号应为：")
table(doc,["位置","左侧","右侧"],[["前轮","M1：正转（+）","M2：反转（−）"],["后轮","M3：反转（−）","M4：正转（+）"]],[1440,3960,3960])
p(doc, "向左横移时，上述四个符号全部反转。负号只表示旋转方向，比较四轮快慢时应比较 RPM 的绝对值。")
callout(doc, "必须先确认", "实际接线中的 M1～M4位置、轮子安装方向和代码定义一致。轮序错误不能靠 PID 修正。", RED)

heading(doc, "3  Keil 中需要观察的变量", 1)
table(doc,["变量","含义","主要用途"],[["target_rpm[0..3]","四轮目标转速","确认运动学计算和目标符号"],["actual_rpm[0..3]","编码器测得的实际转速","判断每个轮子是否跟上目标"],["motor_pwm[0..3]","PID 希望输出的 PWM","判断 PID 是否已到输出上限"],["current_output[0..3]","经过斜坡后的真实 PWM","判断硬件当前实际输出"],["encoder_delta[0..3]","每周期编码器增量","核对编码器方向和是否丢计数"]],[2200,3000,4160])
p(doc, "对每一个轮子，都要同时看 target、actual 和 PWM。只看四个 actual_rpm 是否相同，无法区分目标错误与负载问题。")

heading(doc, "4  三组实验：按顺序执行", 1)
heading(doc, "4.1  实验一：单轮编号与方向检查（底盘架空）", 2)
step(doc,"依次选择 M1～M4","每次只让一个轮子转动，确认它在车体上的实际位置。")
step(doc,"给正目标转速","target_rpm 为正时，actual_rpm 也必须为正；给负目标时，两者也必须同时为负。")
step(doc,"发现符号相反时","检查电机方向、编码器 A/B 相和 encoder_direction，不要继续调 PID。当前方向修正为 {1, −1, 1, −1}。")

heading(doc, "4.2  实验二：架空固定速度横移", 2)
p(doc, "给固定右移指令，使目标转速约为 40～60 RPM。正常示例：")
table(doc,["轮子","目标 RPM","实际 RPM（示例）","判断"],[["M1","+60","+58","正常"],["M2","−60","−59","正常"],["M3","−60","−57","正常"],["M4","+60","+59","正常"]],[1440,1800,2880,3240])
p(doc, "如果悬空时某轮实际转速已经明显偏低，先检查该轮编码器、接线、电机与 PID；不要进入落地比较。")

heading(doc, "4.3  实验三：落地固定速度横移", 2)
p(doc, "保持与实验二相同的目标速度、电池电量和运动方向，让底盘落地。等待起步过程结束后，记录稳定运行阶段的数据。典型异常示例：")
table(doc,["轮子","目标 RPM","实际 RPM","PWM","解释"],[["M1","+60","+57","+100","能跟随"],["M2","−60","−42","−150","明显掉速且输出到顶"],["M3","−60","−55","−110","基本正常"],["M4","+60","+56","+105","基本正常"]],[1200,1440,1440,1200,4080])
callout(doc, "示例结论", "M2 的 PID 已经给到 −150，但实际只能达到 −42 RPM，说明控制器已经尽力。此时重点检查 M2机械阻力、轮子受力和当前 PWM 上限。")

heading(doc, "5  如何根据数据定位原因", 1)
table(doc,["观察结果","最可能原因","下一步"],[["实际 RPM 偏低，PWM 接近 ±150","输出上限不足或机械负载过大","检查卡涩、压地和电源；再小幅提高输出上限验证"],["实际 RPM 偏低，PWM 仍较小","PID增益/反馈异常","核对编码器；逐步调整 Kp、Ki"],["四轮实际 RPM 均能跟随，但车仍斜走","麦轮滚子、车架、重心或打滑","检查四轮压地、滚子和安装方向"],["只有起步时偏，稳定后变直","电机死区或 PWM 斜坡响应不同","比较 current_output，考虑启动补偿或斜坡调整"],["横移路线基本正确，但车头旋转","对角轮合力不平衡，且无航向闭环","先修机械/轮速，再增加陀螺仪航向保持"]],[2640,3120,3600])

heading(doc, "6  本工程中需要重点关注的参数", 1)
table(doc,["项目","当前值","含义"],[["四轮速度 PID","Kp=1.4，Ki=0.05，Kd=0","四轮使用相同参数"],["PID输出范围","−150～+150","PID最大只能请求150 PWM"],["硬件PWM满量程","419","当前PID上限约为满量程的36%"],["控制周期","约10 ms","编码器测速与PID更新周期"],["PWM斜坡","每周期最多变化15","约100 ms从150降至0"],["编码器每圈计数","1404","13 PPR × 4倍频 × 27减速比"],["底盘最大轮速","80 RPM","运动学层的轮速限制"]],[2640,2640,4080])
p(doc, "当前 PID 输出上限 150，而硬件满量程为 419。悬空时 150 可能足够，落地横移时则可能无法提供足够力矩。这是与现象最吻合的首要检查点。")

heading(doc, "7  参数调整建议", 1)
step(doc,"先记录原始数据","在不改参数时完成悬空与落地对照，保留基准。")
step(doc,"先排除机械问题","确认滚子无卡涩、四轮同时压地、车架不晃、重心不过分偏置。")
step(doc,"验证输出上限","如果实际转速偏低且 motor_pwm 长期卡在 ±150，可将四轮 PID 输出上限统一临时提高到 ±200，再重复同一实验。")
step(doc,"再调整 Kp","若 PWM 未到上限但跟随慢，逐步提高 Kp；每次改动幅度保持较小，并观察振荡和噪声。")
step(doc,"最后调整 Ki","使用小 Ki 消除稳定后的剩余误差。若出现慢速来回摆动或积分积累，应减小 Ki。")
callout(doc, "安全提示", "不要一开始把输出上限直接改到 ±419。每次测试保持低速、留出停车空间，并持续观察电流、电机与驱动器温度。", RED)

heading(doc, "8  机械检查清单", 1)
for item in ["底盘放在平整地面，按压四角，确认没有三轮着地或车架晃动。","手动拨动所有麦轮滚子，确认每个滚子都能顺畅转动且阻力接近。","确认四个麦轮安装方向与代码假设的 X 型运动学一致。","检查轮轴是否平行、轮子高度是否一致、紧固件是否松动。","检查电池和载荷是否明显偏向某一侧或某个角。","左右各横移一次：若总向同一车身方向偏，优先查机械/重心；若偏移方向随横移反转，优先查轮速匹配和运动学。"]: bullet(doc,item)

heading(doc, "9  现场数据记录表", 1)
p(doc, "建议在电池电量相近、目标速度相同的条件下填写。RPM比较时看绝对值，同时保留正负号用于核对方向。")
table(doc,["状态/方向","轮子","目标RPM","实际RPM","motor_pwm","current_output","备注"],[["悬空/右移","M1","","","","",""],["","M2","","","","",""],["","M3","","","","",""],["","M4","","","","",""],["落地/右移","M1","","","","",""],["","M2","","","","",""],["","M3","","","","",""],["","M4","","","","",""]],[1680,840,1320,1320,1440,1680,1080])

heading(doc, "10  最终判断口诀", 1)
callout(doc, "轮速跟不上", "查输出上限、PID、电机、编码器和机械阻力。")
callout(doc, "轮速能跟上但车辆仍偏", "查麦轮安装、滚子、四轮受力、重心和地面打滑。")
callout(doc, "车头旋转", "先消除轮速与机械差异，再考虑增加陀螺仪航向闭环。")

doc.core_properties.title = "麦克纳姆底盘横移偏移排查手册"
doc.core_properties.subject = "Four_Motor_PID_Test 工程现场调试指南"
doc.core_properties.author = "Codex"
doc.save(OUT)
print(OUT)
