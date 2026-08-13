from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "RoboGame2026从远程开发到现场推进_大白话总结.md"
OUTPUT = ROOT / "RoboGame2026从远程开发到现场推进_大白话总结.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "17365D"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
CODE_FILL = "F2F4F7"
RED = "9B1C1C"
GOLD = "7A5A00"
WHITE = "FFFFFF"
BLACK = "202124"
FONT_LATIN = "Calibri"
FONT_CJK = "Microsoft YaHei"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    total = sum(widths_dxa)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, *, size=None, bold=None, italic=None, color=BLACK, font=FONT_LATIN) -> None:
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def restarted_numbering_id(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    style = doc.styles["List Number"]
    num_id = style._element.pPr.numPr.numId.val
    source_num = numbering.xpath(f'./w:num[@w:numId="{num_id}"]')[0]
    new_num_id = max(int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))) + 1
    new_num = deepcopy(source_num)
    new_num.set(qn("w:numId"), str(new_num_id))
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    new_num.append(override)
    numbering.append(new_num)
    return new_num_id


def apply_numbering_id(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = num_id


def add_inline_markdown(paragraph, text: str, *, base_size=11, base_color=BLACK) -> None:
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            set_run_font(paragraph.add_run(text[pos:match.start()]), size=base_size, color=base_color)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=10, color=DARK_BLUE, font="Consolas")
            run.font.highlight_color = None
        else:
            set_run_font(paragraph.add_run(token[2:-2]), size=base_size, bold=True, color=base_color)
        pos = match.end()
    if pos < len(text):
        set_run_font(paragraph.add_run(text[pos:]), size=base_size, color=base_color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(paragraph.add_run("第 "), size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=MUTED)
    set_run_font(paragraph.add_run(" 页"), size=9, color=MUTED)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(hp.add_run("RoboGame2026 · 现场推进参考手册"), size=8.5, bold=True, color=MUTED)
    footer = section.footer
    add_page_number(footer.paragraphs[0])


def add_cover(doc: Document) -> None:
    for _ in range(5):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(kicker.add_run("ROBOGAME 2026"), size=11, bold=True, color=BLUE)
    kicker.paragraph_format.space_after = Pt(18)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    set_run_font(title.add_run("从远程开发到现场推进"), size=28, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    set_run_font(subtitle.add_run("大白话项目总结与现场行动手册"), size=16, color=DARK_BLUE)

    rule = doc.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(rule.add_run("━━━━━━━━━━━━━━━━━━━━"), size=10, color=LIGHT_BLUE)
    rule.paragraph_format.space_after = Pt(28)

    lead = doc.add_paragraph()
    lead.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead.paragraph_format.left_indent = Inches(0.8)
    lead.paragraph_format.right_indent = Inches(0.8)
    lead.paragraph_format.space_after = Pt(12)
    set_run_font(
        lead.add_run("给需要同时推进项目、理解代码并组织现场协作的低年级本科生"),
        size=11.5,
        color=MUTED,
    )

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(meta.add_run("版本：2026-08-08  ·  当前软件检查点：304623d"), size=10, color=MUTED)
    doc.add_page_break()

    p = doc.add_paragraph(style="Heading 1")
    p.add_run("如何使用这份手册")
    for item in (
        "第一次阅读：先看第一、三、四部分，理解项目现状、现场顺序和异常处理。",
        "每天到场前：查看第五、六、七部分，明确会议节奏、记录模板和停止条件。",
        "出现问题时：不要从头翻阅，直接在第四部分按现象寻找应对办法。",
        "学习复盘时：用每一节的“为什么”解释设计，而不是只记命令。",
    ):
        para = doc.add_paragraph(style="List Bullet")
        add_inline_markdown(para, item)

    callout = doc.add_table(rows=1, cols=1)
    set_table_geometry(callout, [9360])
    cell = callout.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE)
    set_cell_margins(cell, 160, 180, 160, 180)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_after = Pt(0)
    set_run_font(cp.add_run("核心原则：先确认信息是真的，再让动作可控；先证明一定能停，再尝试完整闭环。"), size=11.5, bold=True, color=NAVY)

    doc.add_paragraph()
    contents = doc.add_paragraph(style="Heading 1")
    contents.add_run("内容导航")
    sections = [
        "第一部分：从开始到现在，我们做了什么",
        "第二部分：为什么一直强调小步推进",
        "第三部分：到现场以后具体要干什么",
        "第四部分：现场可能遇到什么问题，以及怎么办",
        "第五部分：怎么在开会中保持节奏",
        "第六部分：现场每日工作模板",
        "第七部分：什么时候必须停止测试",
        "第八部分：怎样判断项目真的完成了一个闭环",
    ]
    for idx, text in enumerate(sections, 1):
        para = doc.add_paragraph(style="List Number")
        add_inline_markdown(para, text)
    doc.add_page_break()


def widths_for_table(rows: list[list[str]]) -> list[int]:
    cols = max(len(row) for row in rows)
    if cols == 2:
        first = max(len(row[0]) if row else 0 for row in rows)
        return [2700, 6660] if first > 12 else [2100, 7260]
    if cols == 3:
        return [1700, 3830, 3830]
    if cols == 4:
        return [2100, 2100, 2580, 2580]
    if cols == 5:
        return [1600, 1900, 1900, 1980, 1980]
    base = 9360 // cols
    widths = [base] * cols
    widths[-1] += 9360 - sum(widths)
    return widths


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    widths = widths_for_table(rows)
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if r_idx == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            text = row[c_idx] if c_idx < len(row) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            add_inline_markdown(p, text, base_size=9.4 if cols >= 4 else 10, base_color=NAVY if r_idx == 0 else BLACK)
            for run in p.runs:
                if r_idx == 0:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code_block(doc: Document, lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CODE_FILL)
    set_cell_margins(cell, 120, 180, 120, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    for idx, line in enumerate(lines):
        if idx:
            p.add_run().add_break()
        set_run_font(p.add_run(line), size=9.2, color=DARK_BLUE, font="Consolas")


def add_quote(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "FFF8E8")
    set_cell_margins(cell, 130, 200, 130, 200)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_inline_markdown(p, text, base_size=11, base_color=GOLD)
    for run in p.runs:
        run.italic = True


def parse_markdown(doc: Document, text: str) -> None:
    lines = text.splitlines()
    i = 0
    in_code = False
    code_lines: list[str] = []
    skipped_source_title = False
    active_numbering_id: int | None = None
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        if line.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not line or line == "---":
            active_numbering_id = None
            i += 1
            continue
        if line.startswith("# ") and not skipped_source_title:
            skipped_source_title = True
            i += 1
            while i < len(lines) and (not lines[i].strip() or "生成日期" in lines[i] or "适合阅读" in lines[i] or "用途" in lines[i]):
                i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Heading 1")
            p.paragraph_format.page_break_before = True
            add_inline_markdown(p, line[2:].strip(), base_size=16, base_color=BLUE)
            active_numbering_id = None
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline_markdown(p, line[3:].strip(), base_size=13, base_color=BLUE)
            active_numbering_id = None
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 3")
            add_inline_markdown(p, line[4:].strip(), base_size=12, base_color=DARK_BLUE)
            active_numbering_id = None
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+", lines[i + 1]):
            active_numbering_id = None
            table_rows: list[list[str]] = []
            table_rows.append([cell.strip() for cell in line.strip("|").split("|")])
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_rows.append([cell.strip() for cell in lines[i].strip().strip("|").split("|")])
                i += 1
            add_markdown_table(doc, table_rows)
            continue
        if line.startswith("> "):
            active_numbering_id = None
            quote_lines = [line[2:].strip()]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("> "):
                quote_lines.append(lines[i].strip()[2:].strip())
                i += 1
            add_quote(doc, " ".join(quote_lines))
            continue
        checkbox = re.match(r"^- \[([ xX])\] (.*)$", line)
        if checkbox:
            active_numbering_id = None
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.375)
            p.paragraph_format.first_line_indent = Inches(-0.188)
            p.paragraph_format.space_after = Pt(4)
            marker = "☒" if checkbox.group(1).lower() == "x" else "☐"
            set_run_font(p.add_run(f"{marker} "), size=11, color=BLUE)
            add_inline_markdown(p, checkbox.group(2))
            i += 1
            continue
        bullet = re.match(r"^- (.*)$", line)
        if bullet:
            active_numbering_id = None
            p = doc.add_paragraph(style="List Bullet")
            add_inline_markdown(p, bullet.group(1))
            i += 1
            continue
        number = re.match(r"^\d+\. (.*)$", line)
        if number:
            if active_numbering_id is None:
                active_numbering_id = restarted_numbering_id(doc)
            p = doc.add_paragraph(style="List Number")
            apply_numbering_id(p, active_numbering_id)
            add_inline_markdown(p, number.group(1))
            i += 1
            continue
        active_numbering_id = None
        p = doc.add_paragraph()
        add_inline_markdown(p, line)
        i += 1


def build() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    parse_markdown(doc, text)
    core = doc.core_properties
    core.title = "RoboGame2026：从远程开发到现场推进"
    core.subject = "大白话项目总结与现场行动手册"
    core.author = "RoboGame2026 算法组"
    core.keywords = "RoboGame2026, ROS2, 现场联调, 安全验收, 项目管理"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
