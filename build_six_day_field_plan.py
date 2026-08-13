from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "RoboGame2026_六天现场任务计划_2026-08-09至08-14.docx"

BLUE = "2E74B5"
NAVY = "17365D"
DARK = "1F4D78"
MUTED = "667085"
PALE = "E8EEF5"
GRAY = "F4F6F9"
GOLD = "FFF4CE"
RED = "FDECEC"
WHITE = "FFFFFF"
BLACK = "202124"
FONT = "Microsoft YaHei"


DAYS = [
    {
        "date": "8 月 9 日｜第 1 天",
        "title": "硬件盘点、接口冻结与数据采集",
        "goal": "先得到可信的硬件事实，不追求完整演示。今天只观察、测量、确认和留证。",
        "why": "软件已经能在模拟环境闭环，但真实尺寸、相机位置、串口数据和机构动作仍未知。如果跳过这些事实直接联调，出现问题时很难知道错在硬件、协议还是代码。",
        "tasks": [
            "测量轮径、轮距、车体尺寸、夹爪工作区和三层放置高度；拍摄关键位置照片。",
            "确认相机安装位置、视野与机械臂遮挡关系，记录分辨率、帧率和接口。",
            "向电控确认 ODOM、IMU、STATUS 的逐字节字段表、单位、方向和值域，并保存真实十六进制样例。",
            "向机械确认 GRAB、LIFT、RELEASE、STOP 的命令、完成证据、错误码、超时与 CANCEL 行为。",
            "只验证串口设备、握手、RobotStatus、急停和物理启动；不发送自动运动命令。",
            "采集橙色、紫色、空场景及机械臂遮挡照片/短视频，并按日期备份。",
        ],
        "gate": [
            "接口负责人、字段单位、坐标正方向和错误行为都有记录。",
            "关键尺寸有数字，相机能否看见抓取区和放置区已有结论。",
            "原始串口样例、照片和视频至少有两份副本。",
        ],
        "stop": "任何协议仍靠猜测；急停状态异常；相机或机构位置可能碰撞。",
        "fallback": "整理接口问题清单、给固定十六进制样例写解码测试、归档视频并更新 FREEZE_TABLE。",
    },
    {
        "date": "8 月 10 日｜第 2 天",
        "title": "真实状态、里程计与底盘低速安全",
        "goal": "让软件只在可信状态下运动，并证明小车一定能安全停下。",
        "why": "底盘能走并不等于可控。方向、单位、状态新鲜度、急停和失联停车任何一项错误，都会让后续闭环变得危险。",
        "tasks": [
            "按冻结协议接入 STATUS、ODOM、IMU 解码，并做长度、值域、序号和时间戳检查。",
            "确认真实编码器里程计，不用上位机命令速度积分冒充真实运动。",
            "架空或清空场地，逐项验证 vx、vy、wz 的正负方向与单位；一次只改变一个变量。",
            "先测试 STOP，再用极低限速测试前后、左右和旋转；旁边必须有人手持急停。",
            "测试物理急停、拔线失联、状态过期和机构故障时的速度门控。",
            "记录状态频率、延迟、丢帧、低速停车距离及异常日志。",
        ],
        "gate": [
            "六个运动方向与约定一致，速度量级合理。",
            "STOP、急停和通信失联都能让车辆停止。",
            "状态无法解码、过期或复位时不会被报告为健康。",
        ],
        "stop": "STOP 无效、拔线后仍运动、方向与预期不一致、状态频率严重波动。",
        "fallback": "保存原始帧和 rosbag/日志，离线分析解码；只修复阻断安全门槛的问题。",
    },
    {
        "date": "8 月 11 日｜第 3 天",
        "title": "机构单动作真实接入",
        "goal": "独立验收每个机构动作，不运行完整抓放流程。",
        "why": "完整流程失败时变量太多。先证明每个动作能单独成功、失败和被取消，才能知道闭环中的问题属于哪一步。",
        "tasks": [
            "将真实机构协议接入 robot_bridge，保持 mock 与 field 配置严格分离。",
            "用 mechanism_acceptance.py 依次测试 STOP、GRAB、RELEASE、LIFT。",
            "先空载、再带载；每次记录请求、耗时、响应、RobotStatus 和 CSV。",
            "每个动作至少重复 3 次，记录正常值、最大值和明显抖动。",
            "验证超时、限位、故障、急停和人为 CANCEL；观察执行器是否真的停止。",
            "把实测耗时回填 TIME_BUDGET，重新确定安全超时。",
        ],
        "gate": [
            "每个动作都能独立得到明确的成功或失败结果。",
            "STOP/CANCEL 的物理效果得到确认，而非客户端仅停止等待。",
            "CSV 中有动作前后状态、耗时、错误码和说明。",
        ],
        "stop": "动作方向错误、机构撞限位、取消后仍继续运动、状态与实物不一致。",
        "fallback": "机械未就绪时，审核协议样例、完善适配器测试、演练故障注入卡片。",
    },
    {
        "date": "8 月 12 日｜第 4 天",
        "title": "相机标定、抓取证据与放置证据",
        "goal": "把真实视觉数据接入抓取和放置判断，明确遮挡时应该返回什么。",
        "why": "模拟检测不能说明现场光照、视角和遮挡可用。视觉必须区分“稳定”“失败”和“看不清”，不能把看不到误当成成功。",
        "tasks": [
            "固定相机安装姿态，记录分辨率、帧率、曝光、ROI 以及相机到夹爪的偏移。",
            "采集橙色、紫色、空场景、反光、近远距离和机械臂遮挡视频。",
            "用本地标注网页处理新视频，检查时间线，补充时间段标注并生成 HTML 报告。",
            "按一次只改一组参数的原则调整 HSV、ROI、面积和连续帧门槛。",
            "比较机械臂撤退前后视野，确定稳定 3 秒从何时开始；保留允许短暂丢帧的可调间隔。",
            "与机械确认抓取成功证据：视觉、夹爪状态、负载传感器或组合策略。",
        ],
        "gate": [
            "真实视频能区分目标与空场景，报告可回看。",
            "遮挡或证据不足时返回 INCONCLUSIVE，而不是伪造 STABLE。",
            "相机安装方案和 3 秒观察起点有书面结论。",
        ],
        "stop": "相机视野被长期遮挡、曝光不稳定、检测结果与肉眼明显相反。",
        "fallback": "不能拍新视频时，处理已有数据、完善报告索引和参数对比，不现场盲调。",
    },
    {
        "date": "8 月 13 日｜第 5 天",
        "title": "二审主链路与第一个单方块物理闭环",
        "goal": "从启动区出发，完成到取料位置、取出并在本体内保持、运送、放置、撤退和稳定观察。",
        "why": "二审把“从启动区到存矿位置”“取出并存放于机器人本体”和“正常搭建”分成三个 20 分项。必须按评分动作拍到完整证据，不能只在夹爪旁演示抓放。",
        "tasks": [
            "分别复验接近、对准、GRAB、LIFT、RELEASE、RETREAT 和稳定观察。",
            "使用现场实测的动作超时，不沿用未经验证的旧假设；全程低速。",
            "从 600 mm × 600 mm 启动区内开始，运行一次完整流程；急停负责人站在安全位置并只负责安全。",
            "单独证明方块被取出后能在机器人本体中稳定保持/携带，再进入运送和放置阶段。",
            "保存终端日志、ROS 日志、状态、视频、检测 JSONL/HTML 和参数快照。",
            "失败时归类为视觉、定位、运动、通信、机构、证据、超时或人为取消。",
            "一次只修一个阻断闭环的问题；修复后从相关最小步骤重新验证。",
        ],
        "gate": [
            "至少一次从启动区开始的完整闭环有连续、可追溯的证据。",
            "取料后方块确实由机器人本体保持，不依靠人员或场外辅助。",
            "失败时能够安全停止，并指出失败阶段和直接证据。",
            "没有为了跑通而关闭安全门控或伪造状态。",
        ],
        "stop": "任何安全门槛未通过、无法解释的自主运动、证据链中断或现场人员意见不一致。",
        "fallback": "退回最后一个已通过阶段，重跑单动作/单传感器验收，不新增架构。",
    },
    {
        "date": "8 月 14 日｜第 6 天",
        "title": "重复性、二审取证、故障恢复与交接",
        "goal": "证明系统可以重复，并按二审评分项形成可提交的视频证据。",
        "why": "二审不仅看代码是否正确，还按视频中的具体动作评分。需要把技术验收结果转换成评委一眼能看懂的镜头。",
        "tasks": [
            "在相同条件下重复单方块流程，统计成功率、耗时和主要失败原因。",
            "若单层已稳定，尝试第二层搭建并保持至少 3 秒；若不稳定，保留单层真实证据并明确后续计划。",
            "复验急停、失联、状态过期、动作超时、CANCEL、MCU 重启和恢复。",
            "按评分项分别拍摄：电控布线、底盘运动、启动区到取料区、取出并本体保持、建筑搭建、完整外观、急停断电。",
            "急停必须单独拍摄：机器人执行器正在运行 → 按下显著急停 → 所有执行器停止/断电。",
            "冻结可用参数与对应硬件版本，不再临时同时修改多个参数。",
            "整理一键启动/停止步骤、已知限制、恢复方法和现场注意事项。",
            "备份视频、CSV、JSONL、HTML、ROS 日志、配置和代码提交号。",
            "更新 TODO、学习记录和交接说明；创建可回滚的 Git 检查点。",
        ],
        "gate": [
            "有明确的重复次数、成功率、平均/最大耗时和失败分类。",
            "七类二审证据镜头均有独立文件或时间段索引，并能对应评分表。",
            "重启与故障后能按文档恢复，队友可以照步骤复现。",
            "证据目录完整，参数、硬件版本和代码提交相互对应。",
        ],
        "stop": "为了提高成功率临时隐藏失败、删除异常数据或跳过安全复验。",
        "fallback": "闭环仍不稳定时，交付可重复的最小能力和明确阻塞项，不包装成已完成。",
    },
]


def font(run, size=10.5, bold=False, color=BLACK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    shd = pr.find(qn("w:shd")) or OxmlElement("w:shd")
    if shd.getparent() is None:
        pr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, v=90, h=120):
    pr = cell._tc.get_or_add_tcPr()
    mar = pr.find(qn("w:tcMar")) or OxmlElement("w:tcMar")
    if mar.getparent() is None:
        pr.append(mar)
    for name, val in (("top", v), ("bottom", v), ("start", h), ("end", h)):
        el = mar.find(qn(f"w:{name}")) or OxmlElement(f"w:{name}")
        if el.getparent() is None:
            mar.append(el)
        el.set(qn("w:w"), str(val)); el.set(qn("w:type"), "dxa")


def repeat_header(row):
    pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); pr.append(el)


def fixed_table(table, widths):
    table.autofit = False
    pr = table._tbl.tblPr
    total = sum(widths)
    for tag, value in (("tblW", total), ("tblInd", 120)):
        el = pr.find(qn(f"w:{tag}")) or OxmlElement(f"w:{tag}")
        if el.getparent() is None: pr.append(el)
        el.set(qn("w:w"), str(value)); el.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        el = OxmlElement("w:gridCol"); el.set(qn("w:w"), str(width)); grid.append(el)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            margins(cell)
            tcw = cell._tc.get_or_add_tcPr().find(qn("w:tcW")) or OxmlElement("w:tcW")
            if tcw.getparent() is None: cell._tc.get_or_add_tcPr().append(tcw)
            tcw.set(qn("w:w"), str(widths[i])); tcw.set(qn("w:type"), "dxa")


def add_page_number(p):
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(p.add_run("第 "), 8.5, color=MUTED)
    r = p.add_run(); begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText"); text.set(qn("xml:space"), "preserve"); text.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    r._r.extend([begin, text, end]); font(r, 8.5, color=MUTED)
    font(p.add_run(" 页"), 8.5, color=MUTED)


def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.72)
    sec.left_margin = sec.right_margin = Inches(0.85)
    sec.header_distance, sec.footer_distance = Inches(0.32), Inches(0.32)
    normal = doc.styles["Normal"]
    normal.font.name = FONT; normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5); normal.paragraph_format.space_after = Pt(4); normal.paragraph_format.line_spacing = 1.18
    for name, size, color, before, after in (("Heading 1", 18, BLUE, 12, 7), ("Heading 2", 13, DARK, 9, 5), ("Heading 3", 11.5, DARK, 7, 4)):
        s = doc.styles[name]; s.font.name = FONT; s._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after); s.paragraph_format.keep_with_next = True
    header = sec.header.paragraphs[0]; font(header.add_run("RoboGame2026｜六天现场任务计划（第一版）"), 8.5, True, MUTED)
    add_page_number(sec.footer.paragraphs[0])


def para(doc, text, size=10.5, bold=False, color=BLACK, after=4, align=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    if align is not None: p.alignment = align
    font(p.add_run(text), size, bold, color)
    return p


def callout(doc, label, text, fill=PALE):
    t = doc.add_table(rows=1, cols=1); fixed_table(t, [9360]); shade(t.cell(0, 0), fill)
    p = t.cell(0, 0).paragraphs[0]; font(p.add_run(label + "："), 10.2, True, DARK); font(p.add_run(text), 10.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def checklist(doc, items):
    for item in items:
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.12); p.paragraph_format.first_line_indent = Inches(-0.02); p.paragraph_format.space_after = Pt(2.4)
        font(p.add_run("□  "), 10.5, True, BLUE); font(p.add_run(item), 9.8)


def blank_lines(doc, labels):
    for label in labels:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8)
        font(p.add_run(label + "："), 9.8, True, DARK); font(p.add_run("____________________________________________________________"), 9.8, color=MUTED)


def day_page(doc, day):
    start = para(doc, day["date"], 10, True, BLUE, 2)
    start.paragraph_format.page_break_before = True
    para(doc, day["title"], 20, True, NAVY, 9)
    callout(doc, "今日唯一目标", day["goal"], PALE)
    callout(doc, "为什么今天做这些", day["why"], GRAY)
    doc.add_heading("按顺序执行", level=2)
    checklist(doc, day["tasks"])
    doc.add_heading("进入下一天前必须通过", level=2)
    checklist(doc, day["gate"])
    callout(doc, "立即停止条件", day["stop"], RED)
    callout(doc, "硬件未就绪时的替代任务", day["fallback"], GOLD)
    record = para(doc, day["date"] + "｜现场记录页", 17, True, NAVY, 7)
    record.paragraph_format.page_break_before = True
    blank_lines(doc, ["今日负责人", "急停负责人", "开始时间", "结束时间", "代码提交号", "硬件/固件版本"])
    doc.add_heading("验收记录", level=2)
    table = doc.add_table(rows=1, cols=4); fixed_table(table, [2350, 1200, 2900, 2910]); repeat_header(table.rows[0])
    for i, text in enumerate(("检查项", "结果", "证据路径/编号", "问题与下一步")):
        shade(table.cell(0, i), PALE); p = table.cell(0, i).paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run(text), 9.2, True, DARK)
    for item in day["gate"]:
        cells = table.add_row().cells
        values = (item, "□通过\n□未通过", "", "")
        for i, text in enumerate(values): font(cells[i].paragraphs[0].add_run(text), 8.8)
    fixed_table(table, [2350, 1200, 2900, 2910])
    doc.add_heading("当天收尾（必须填写）", level=2)
    blank_lines(doc, ["今天完成了什么", "没有完成的原因", "新增问题/风险", "修改的参数或代码", "数据保存路径", "明天第一步"])
    para(doc, "现场原则：没有证据就不算通过；无法判断时写 INCONCLUSIVE，不写 PASS。", 9.8, True, BLUE, 0)


def add_overview(doc):
    para(doc, "现场执行手册｜第一版参考", 10.5, True, BLUE, 4, WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "RoboGame2026", 30, True, NAVY, 4, WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "六天现场任务计划", 22, True, DARK, 8, WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "2026 年 8 月 9 日—8 月 14 日", 13, True, MUTED, 16, WD_ALIGN_PARAGRAPH.CENTER)
    callout(doc, "总目标", "在不跳过安全阶段的前提下，完成并重复验证一个单方块物理闭环：真实状态可信 → 底盘可控且可停 → 单机构动作可控且可取消 → 真实视觉证据 → PICK/PLACE → 撤退 → 稳定观察 → 可重复与可恢复。", PALE)
    para(doc, "使用方法", 14, True, NAVY, 6)
    checklist(doc, [
        "每天开工前只读当天两页，不提前做后一天的危险动作。",
        "每完成一项就在方框内打勾，并把证据路径写在记录页。",
        "验收门槛未全部通过时，不进入下一阶段；先执行替代任务。",
        "每天结束前填写六项收尾记录，并备份数据。",
        "现场事实与本计划冲突时，以安全和真实测量为准，并记录变更原因。",
    ])
    para(doc, "六天节奏总览", 14, True, NAVY, 6)
    table = doc.add_table(rows=1, cols=3); fixed_table(table, [1450, 2950, 4960]); repeat_header(table.rows[0])
    for i, text in enumerate(("日期", "主任务", "当天输出")):
        shade(table.cell(0, i), PALE); font(table.cell(0, i).paragraphs[0].add_run(text), 9.4, True, DARK)
    outputs = ["冻结表、硬件尺寸、协议样例、现场视频", "状态/方向/停车/失联安全证据", "动作 CSV、取消证据、实测时间预算", "标注时间线、HTML 报告、证据策略", "一次完整闭环及失败分类", "重复性统计、恢复流程、交接包"]
    for d, out in zip(DAYS, outputs):
        cells = table.add_row().cells
        for i, text in enumerate((d["date"].split("｜")[0], d["title"], out)): font(cells[i].paragraphs[0].add_run(text), 8.15)
    fixed_table(table, [1450, 2950, 4960])
    doc.add_page_break()
    para(doc, "全程不可违反的规则", 19, True, NAVY, 8)
    checklist(doc, [
        "STOP、急停和失联停车未通过前，不运行自动导航或完整机构流程。",
        "未经电控/机械确认，不猜测协议字节、动作方向、限位和负载。",
        "每次只引入一个真实变量，每次只改一组参数。",
        "运动测试至少两人：一人操作，一人只负责急停和观察危险。",
        "从 8 月 13 日起停止非必要重构，只做闭环阻塞修复、回归和恢复。",
        "大视频和 results 不直接提交普通 Git；重要数据必须外部备份。",
    ])
    callout(doc, "会议节奏", "每天用 10 分钟说明：唯一目标、昨天证据、今天分工、进入下一阶段的门槛、最大风险。遇到争论先回到数据和接口定义，不在现场同时尝试多个改法。", GOLD)
    para(doc, "统一问题分类", 14, True, NAVY, 6)
    table = doc.add_table(rows=1, cols=3); fixed_table(table, [1800, 3640, 3920]); repeat_header(table.rows[0])
    rows = [
        ("通信", "串口、握手、心跳、丢帧、状态过期", "保存原始帧和日志，先停运动"),
        ("运动", "方向、单位、限速、停车距离", "回到单轴低速测试"),
        ("机构", "超时、限位、故障、CANCEL", "回到单动作空载验收"),
        ("视觉", "光照、ROI、遮挡、误检/漏检", "保存视频，离线报告后再调参"),
        ("证据", "看到动作但无法证明结果", "返回 INCONCLUSIVE，补证据来源"),
        ("流程", "步骤顺序、状态机、恢复入口", "退回最后一个已通过阶段"),
    ]
    for i, text in enumerate(("类别", "典型现象", "第一反应")):
        shade(table.cell(0, i), PALE); font(table.cell(0, i).paragraphs[0].add_run(text), 9.3, True, DARK)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row): font(cells[i].paragraphs[0].add_run(text), 8.9)
    fixed_table(table, [1800, 3640, 3920])


def add_review_audit(doc):
    title = para(doc, "二审要求反向审计", 20, True, NAVY, 5)
    title.paragraph_format.page_break_before = True
    para(doc, "依据：《二审评分细则（竞技组）》与《RoboGame2026 竞技组规则手册 1.1》。结论：六天计划的技术顺序合理，但原版只保证“单方块闭环”，不足以自动等同于二审满分。下面把评分证据正式纳入第 5—6 天。", 10.4, False, BLACK, 8)
    table = doc.add_table(rows=1, cols=5); fixed_table(table, [950, 2500, 1760, 2550, 1600]); repeat_header(table.rows[0])
    headers = ("分值", "评分项", "六天计划覆盖", "必须拍到的证据", "风险判断")
    for i, text in enumerate(headers):
        shade(table.cell(0, i), PALE); font(table.cell(0, i).paragraphs[0].add_run(text), 8.7, True, DARK)
    rows = [
        ("10", "电控布线", "第 1 天盘点；第 6 天取证", "整洁固定、绝缘、防松、接口标识", "依赖电控整理"),
        ("10", "底盘运动", "第 2 天安全运动", "前后/横移/旋转，动作稳定", "较可控"),
        ("20", "启动区运行至存矿/取料位置", "第 2 天底盘＋第 5 天主链路", "从启动区完整出发并到达目标位置", "依赖定位/决策"),
        ("20", "取出并存放于机器人本体", "第 3 天机构＋第 5 天闭环", "取出后在本体内稳定保持/携带", "需确认机构定义"),
        ("20", "建筑搭建", "第 4 天证据＋第 5/6 天放置", "方块进入己方搭建区且结构成立；争取第二层", "单层可能说服力弱"),
        ("10", "形态完整", "第 1 天盘点＋第 6 天取证", "全车外观、机构完整、无明显临时散乱", "依赖机械收尾"),
        ("5", "急停开关可控", "第 2 天安全测试＋第 6 天单拍", "运行中按急停，所有执行器停止/断电", "必须独立视频"),
        ("5", "整体运行附加分", "第 6 天最终视频", "节奏流畅、少停顿、画面完整", "有余力再追求"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row): font(cells[i].paragraphs[0].add_run(text), 8.15)
    fixed_table(table, [950, 2500, 1760, 2550, 1600])
    callout(doc, "现实判断", "如果机械、电控和场地都在 8 月 9 日就绪，六天足够冲击“可展示的二审主链路”；但不能现在承诺满分。最不确定的是电控布线/形态完整（不是纯算法任务）、本体存矿机构是否已完成，以及建筑是否能稳定达到第二层。", GOLD)
    callout(doc, "最低交付", "至少形成：安全可控底盘 + 从启动区到目标位置 + 单方块取出并本体保持 + 放入搭建区 + 独立急停视频。即使高层建筑未完成，也能覆盖多数核心评分动作并真实展示当前完成度。", PALE)
    doc.add_heading("二审视频拍摄清单", level=2)
    checklist(doc, [
        "镜头 1：全车静态环绕，展示形态完整和关键机构。",
        "镜头 2：电控舱近景，展示固定、绝缘、走线和急停标识。",
        "镜头 3：底盘前后、横移、旋转，画面中保留整车。",
        "镜头 4：机器人完全位于启动区内，随后自主到达取料位置。",
        "镜头 5：取出方块，并清楚展示方块由机器人本体稳定保持。",
        "镜头 6：运送至搭建区并完成搭建；镜头持续到结构稳定至少 3 秒。",
        "镜头 7（必须单独拍）：执行器运行时按下急停，所有执行器立即停止/断电。",
        "每个原始视频保留原文件；另建一份剪辑/提交副本，不覆盖原始证据。",
    ])


def build():
    doc = Document(); setup(doc); add_overview(doc); add_review_audit(doc)
    for day in DAYS: day_page(doc, day)
    final_title = para(doc, "最终交接检查表", 20, True, NAVY, 8)
    final_title.paragraph_format.page_break_before = True
    checklist(doc, [
        "代码提交号、分支、Ubuntu 构建结果和启动命令已记录。",
        "硬件版本、固件版本、串口设备名和协议版本已记录。",
        "安全测试、单动作、视觉报告、闭环视频和重复性统计均可定位。",
        "可用参数及其适用硬件条件已冻结。",
        "已知限制、未解决问题、恢复步骤和下一位负责人明确。",
        "关键数据至少有本机与外部存储两份副本。",
    ])
    doc.add_heading("最终结论（现场填写）", level=2)
    blank_lines(doc, ["当前可重复能力", "成功率与样本数", "最主要失败原因", "尚未开放的能力", "下一阶段第一任务", "交接人/接收人/日期"])
    callout(doc, "合格的结论", "允许写“目前只稳定完成到单动作/单方块第一层”。不允许把一次偶然成功写成系统已经完成。真实、可复现、能恢复，比演示得好看更重要。", PALE)
    doc.core_properties.title = "RoboGame2026 六天现场任务计划（2026-08-09 至 08-14）"
    doc.core_properties.subject = "现场接入、安全验收、单方块闭环与交接"
    doc.core_properties.author = "RoboGame2026 团队"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
