from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "docs" / "视觉模块项目开发与代码学习总结.docx"
BLUE = "2E74B5"
DARK = "1F4D78"
LIGHT = "F2F4F7"
PALE_BLUE = "EAF2F8"
PALE_GREEN = "EAF5EE"
MUTED = "5B6573"


def font(run, size=11, bold=False, color="000000", italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[i])
            margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    mark_header_row(table.rows[0])
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, LIGHT)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(p.add_run(value), 10, True, DARK)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            font(p.add_run(value), 9.5)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.5 if level == 0 else 0.75)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    font(p.add_run(text), 11)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(6)
    font(p.add_run(text), 11)
    return p


def add_callout(doc, label, text, fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade(cell, fill)
    margins(cell, 150, 150, 180, 180)
    p = cell.paragraphs[0]
    font(p.add_run(label + "  "), 11, True, DARK)
    font(p.add_run(text), 11)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance, sec.footer_distance = Inches(0.45), Inches(0.45)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10
for name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK, 8, 4),
]:
    st = styles[name]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)

# Running furniture
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(hp.add_run("RoboGame 2026  ·  视觉模块学习复盘"), 8.5, color=MUTED)
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(fp.add_run("从少量编程基础走向可验证的工程开发"), 8.5, color=MUTED)

# Editorial cover
for _ in range(4):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("项目学习总结"), 11, True, BLUE)
p.paragraph_format.space_after = Pt(16)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("通过视觉模块，我们学到了什么"), 26, True, DARK)
p.paragraph_format.space_after = Pt(10)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("面向只有少量编程基础的低年级本科生"), 14, color=MUTED)
p.paragraph_format.space_after = Pt(34)
add_callout(doc, "核心结论", "视觉模块让我们第一次完整经历了“需求拆解—算法实现—参数调试—离线测试—结果记录—ROS2 接入—验收”的工程闭环。学到的不只是 OpenCV 函数，而是如何把一段能运行的代码逐步变成一个可复用、可解释、可验证的项目模块。", PALE_GREEN)
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("基于当前仓库的 cube_perception、实验课、配置与验收工具整理"), 10, italic=True, color=MUTED)
doc.add_page_break()

add_heading(doc, "一、我们实际完成了什么", 1)
p = doc.add_paragraph("视觉模块的任务，是从相机、图片或录像中寻找橙色和紫色方块，并把“看见了什么”转换为其他模块可以使用的数据。当前实现已经形成一条清晰的数据链：")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
add_callout(doc, "数据链", "BGR 图像 → HSV 颜色分割 → 形态学去噪 → 轮廓与形状筛选 → 置信度 → 像素位置 → 距离/横向偏差/偏航误差 → 连续帧确认 → JSONL 或 ROS2 /cubes 输出")
add_table(doc, ["组成", "完成的工作", "初学者能理解的意义"], [
    ("检测器", "识别橙色、紫色区域，过滤面积、长宽比、矩形度、实心度等异常目标", "把图像一步步变成结构化结果"),
    ("参数配置", "HSV、ROI、最小面积、置信度、焦距等集中保存在 JSON/ROS2 参数中", "调试参数不必反复改算法代码"),
    ("时序过滤", "连续多帧确认、目标匹配、丢帧处理、指数平滑", "单帧看到不等于可信，稳定输出更重要"),
    ("离线工具", "处理图片/录像、生成标注视频、记录 JSONL、汇总验收报告", "没有真车也能开发和保留证据"),
    ("ROS2 节点", "订阅图像和相机信息，复用同一检测器并发布 /cubes", "纯算法与系统通信可以分开"),
], [1500, 3900, 3960])
add_callout(doc, "项目成熟度", "目前完成的是可独立运行、可离线验证、可接入 ROS2 的视觉软件模块；最终相机、光照、安装角度、真实方块和整车运动仍需实机验收。准确区分“软件可用”和“现场可用”，本身就是工程能力。", PALE_GREEN)

add_heading(doc, "二、项目开发方面的收获", 1)
add_heading(doc, "1. 从任务出发，而不是从代码出发", 2)
doc.add_paragraph("初学者最自然的做法是先打开代码、搜索函数，然后尝试修改。但在这个模块中，我们逐渐学会先问：输入是什么、输出给谁、怎样才算识别正确、失败时怎样留下证据。")
for text in [
    "需求必须可验证：例如不仅说“能识别方块”，还要规定橙色和紫色样本数量、召回率、持续误检、处理帧率和目标移除后的停止时间。",
    "未知条件要显式记录：相机型号、视场角、安装高度、光照和机械遮挡不能靠猜，应进入配置、假设清单和实机验收项。",
    "完成标准要分层：代码能运行、录像效果正常、ROS2 消息能发布、真实机器人能稳定工作，是四个不同阶段。",
]: add_bullet(doc, text)

add_heading(doc, "2. 把大问题拆成可单独验证的小问题", 2)
for text in [
    "先在人工生成的场景中理解颜色掩膜和轮廓，再处理真实录像。",
    "先让单帧检测器正确，再加入连续帧确认和平滑。",
    "先用独立命令行工具输出结果，再封装为 ROS2 节点。",
    "先保存 JSONL 和标注视频，再根据记录生成统计报告。",
]: add_number(doc, text)
doc.add_paragraph("这样的顺序降低了调试难度：当结果异常时，可以判断问题来自颜色阈值、形状条件、距离模型、时序逻辑、录像读写，还是 ROS2 接口，而不是把所有问题混在一起。")

add_heading(doc, "3. 学会用“证据”推进项目", 2)
doc.add_paragraph("工程开发不是凭肉眼说“看起来还行”。本项目保存原始录像、标注视频、逐帧 JSONL 和统计报告，使每次调参都有可回放、可比较的依据。验收集还应与调参素材分开，避免只对熟悉样本表现良好。")

add_heading(doc, "三、代码阅读与使用方面的收获", 1)
add_table(doc, ["代码知识", "在视觉模块中的体现", "可迁移能力"], [
    ("模块与类", "CubeDetector、DetectorConfig、TemporalDetectionFilter 各自负责清晰职责", "阅读项目时先找对象的边界，而不是逐行硬看"),
    ("数据结构", "检测结果包含颜色、置信度、距离、横向偏差、角度、像素坐标和时间戳", "用结构化数据代替散乱变量"),
    ("函数输入输出", "detect(frame) 返回检测结果与掩膜；update(detections) 返回确认后的结果", "通过函数合同理解代码，而非先理解全部实现"),
    ("异常与参数校验", "检查 HSV、ROI、比例、帧率、验收阈值是否合法", "让错误尽早暴露，并给出可读提示"),
    ("命令行参数", "--source、--config、--jsonl、--output-video、--headless 等", "同一份代码可用于预览、批处理和验收"),
    ("文件与序列化", "JSON 保存配置，JSONL 保存逐帧记录，视频保存视觉证据", "学会选择适合的数据格式"),
], [1600, 4000, 3760])
add_heading(doc, "对初学者最重要的代码阅读顺序", 2)
for text in [
    "先读 README，知道模块解决什么问题、怎样运行。",
    "再读配置文件，认识可以调节的参数及其量纲。",
    "看入口文件或 main()，理解程序从哪里开始、数据怎样流动。",
    "看核心类的公开方法，先抓住输入和输出。",
    "最后再进入颜色分割、轮廓计算、评分公式等细节。",
    "每理解一层就运行一个小实验，避免只看不做。",
]: add_number(doc, text)

add_heading(doc, "四、计算机视觉知识的收获", 1)
add_heading(doc, "1. 图像不是“照片”，而是可以计算的数据", 2)
doc.add_paragraph("OpenCV 读取的图像本质上是像素数组。颜色空间、阈值、形态学操作和轮廓计算，都是对数组进行变换。理解这一点后，视觉算法不再神秘：它是在逐步缩小候选区域。")

# Visual example
table = doc.add_table(rows=1, cols=2)
set_table_geometry(table, [4680, 4680])
for cell in table.rows[0].cells:
    margins(cell, 100, 100, 100, 100)
left, right = table.rows[0].cells
left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
shape = left.paragraphs[0].add_run().add_picture(str(ROOT / "lessons" / "02_cube_perception" / "output" / "scene.png"), width=Inches(2.7))
shape._inline.docPr.set("descr", "人工生成的橙色和紫色方块输入场景")
font(left.add_paragraph().add_run("输入场景"), 9, True, MUTED)
right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
shape = right.paragraphs[0].add_run().add_picture(str(ROOT / "lessons" / "02_cube_perception" / "output" / "annotated.png"), width=Inches(2.7))
shape._inline.docPr.set("descr", "视觉检测器输出的方块边界和结果标注图")
font(right.add_paragraph().add_run("检测与标注结果"), 9, True, MUTED)
doc.add_paragraph()

add_heading(doc, "2. HSV、ROI 和形状筛选是互相配合的", 2)
for text in [
    "HSV 阈值负责回答“哪些像素的颜色像目标”。它比直接用 BGR 更适合人工调节，但会受光照、阴影和反光影响。",
    "形态学开闭运算用于去掉小噪点、填补目标内部空洞；核大小过大也可能破坏小目标。",
    "ROI 只在合理区域内搜索，可减少地面、机器人结构和墙面造成的误检，同时降低计算量。",
    "面积、长宽比、矩形度、实心度和最短边等条件共同过滤“颜色相似但形状不对”的区域。",
]: add_bullet(doc, text)

add_heading(doc, "3. 从像素到物理量需要模型与标定", 2)
add_callout(doc, "针孔模型", "距离 ≈ 方块真实宽度 × 像素焦距 ÷ 检测到的像素宽度。横向偏差和偏航角也可由目标中心相对图像中心的位置估计。")
doc.add_paragraph("这让我们认识到：公式只是模型，focal_px 需要用已知距离实测；方块倾斜、遮挡、边界不准都会引入误差。代码中的数字不是天然正确，必须知道来源、单位、适用条件和误差。")

add_heading(doc, "4. 单帧正确不代表系统稳定", 2)
add_table(doc, ["时序参数", "作用", "调得过大或过小的影响"], [
    ("confirm_frames", "连续多少帧后才信任目标", "小：易误报；大：响应变慢"),
    ("match_distance_px", "相邻帧目标允许移动的最大像素距离", "小：目标易断轨；大：可能错误匹配"),
    ("max_missed_frames", "短暂丢失后保留匹配状态的帧数", "小：抗遮挡差；大：旧目标保留过久"),
    ("smoothing_alpha", "新一帧在平滑结果中的权重", "小：更稳但滞后；大：灵敏但抖动"),
], [1800, 3000, 4560])

add_heading(doc, "五、调试、测试和验收方面的收获", 1)
add_heading(doc, "1. 调参是一种受控实验", 2)
for text in [
    "一次尽量只改变一类参数，并记录修改前后的效果。",
    "用亮/暗、远/近、正视/侧视、相邻目标、遮挡和机器人入镜等多种场景测试。",
    "调参素材与最终验收素材分开，防止“记住题目”。",
    "先用 --raw-detections 看单帧检测器，再关闭该选项测试真实的连续帧输出。",
]: add_bullet(doc, text)
add_heading(doc, "2. 正常路径和异常路径都要测试", 2)
doc.add_paragraph("不仅要测试目标清晰可见，还要测试空画面、短暂遮挡、错误颜色、目标移除、视频结束、配置非法和输出路径不可写等情况。真正可靠的程序，要能在失败时停止、给出提示，并避免发布过期目标。")
add_heading(doc, "3. 指标必须知道含义和局限", 2)
add_table(doc, ["指标", "回答的问题", "注意事项"], [
    ("Detection ratio", "已知目标存在的帧中，有多少帧检测到目标？", "只有选定区间每帧都有目标时，才近似召回率"),
    ("Confidence", "当前候选多大程度符合设定的颜色与形状条件？", "不是严格概率，需结合误检分析"),
    ("最长连续检测", "输出是否持续稳定？", "要结合源视频 FPS 换算时间"),
    ("距离误差", "估计距离与实测距离差多少？", "需在多个距离、角度和光照下记录"),
    ("处理帧率", "目标计算平台是否能实时运行？", "录像 FPS、记录频率和纯算法耗时不是同一概念"),
], [1800, 3500, 4060])

add_heading(doc, "4. 视觉验收要按固定流程执行", 2)
doc.add_paragraph("验收不是再看一遍调参录像，而是用未参与调参的材料，按照预先写好的门槛判断模块能否进入底盘联调。当前项目建议采用下面的最低门槛：")
add_table(doc, ["验收项目", "当前建议门槛", "需要保留的证据"], [
    ("样本覆盖", "至少 30 个橙色样例、30 个紫色样例；覆盖亮暗、远近、左右视角、相邻方块和部分遮挡", "原始图片/录像、场景编号和采集条件"),
    ("颜色召回", "橙色、紫色分别达到至少 90%", "逐帧 JSONL、区间说明和统计报告"),
    ("持续误检", "地面、机器人本体、夹爪等位置不得出现持续假目标", "无目标与干扰物录像、标注视频"),
    ("处理能力", "在选定上位机上平均至少 20 FPS", "设备信息、分辨率、源 FPS 与实际处理统计"),
    ("方向正确性", "目标位于左、中、右时，横向偏差方向全部正确", "测试位置表和 /cubes 输出记录"),
    ("距离误差", "在 0.3、0.5、0.8、1.0 m 分别记录，不隐瞒偏差", "卷尺实测值、估计值、绝对/相对误差表"),
    ("目标消失", "移除目标后 0.5 s 内停止报告", "带时间戳的录像和 JSONL"),
], [1700, 4500, 3160])
add_heading(doc, "5. 一次规范验收的六个步骤", 2)
for text in [
    "冻结代码与配置：记录 Git 版本、配置文件、相机、分辨率、帧率和环境。",
    "准备独立验收集：不得继续用同一批素材边验收边调参。",
    "运行独立工具：输出标注视频与 JSONL；正式验收关闭 raw detections，使行为与 ROS2 节点一致。",
    "生成统计报告：按颜色与时间区间计算检测比例、最长连续检测、置信度和距离分布。",
    "人工复核异常：检查持续误检、目标交换、短暂遮挡、移除延迟以及记录时间轴是否正确。",
    "形成结论：通过、失败或证据不足；失败后回到新一轮调参，并更换最终验收数据。",
]: add_number(doc, text)
add_callout(doc, "验收纪律", "验收结果必须能被另一位同学复现。只展示几帧效果图、只报告最好的一段录像，或在验收过程中不断修改阈值，都不能证明模块已经达到联调条件。", PALE_GREEN)

add_heading(doc, "6. 软件验收与实机验收不能混为一谈", 2)
doc.add_paragraph("录像验收通过后，还要在真实安装条件下确认相机视场、畸变、运动模糊、振动、机械臂/夹爪遮挡、场地反光、延迟和时间同步。尤其是放置后稳定观察，如果画面被机构遮住，应输出 INCONCLUSIVE（证据不足），而不是误判为稳定或失败。真实底盘运动必须在视觉门槛通过后再低速接入。")

add_heading(doc, "六、ROS2 与系统集成方面的收获", 1)
add_heading(doc, "1. 纯算法和通信外壳分离", 2)
doc.add_paragraph("CubeDetector 不依赖 ROS2，既可被录像工具调用，也可被 ROS2 节点调用。节点只负责订阅 Image/CameraInfo、图像格式转换、读取参数和发布 CubeDetectionArray。这种结构提高了复用性，也让 Windows 离线测试与 Ubuntu 机器人运行共享同一套核心逻辑。")
add_heading(doc, "2. 接口是团队合同", 2)
doc.add_paragraph("视觉输出不只是“中心点”，而是包含颜色、置信度、距离、横向偏差、偏航误差、像素坐标和时间戳。下游对准与抓取模块依赖这些字段，因此字段含义、单位、坐标正方向和时间有效性必须明确。修改接口会影响多个包，需要重新构建并联调。")
add_heading(doc, "3. 参数化让现场调整不必改代码", 2)
doc.add_paragraph("不同相机和光照可以使用不同配置；比赛日不直接修改共享默认文件，而是复制并保存现场配置。参数化不是为了“看起来高级”，而是为了让算法、设备差异和实验条件彼此分离。")

add_heading(doc, "七、工具与协作方面的收获", 1)
add_table(doc, ["工具/习惯", "本项目中的用途", "形成的能力"], [
    ("Git", "小步提交、保留修改历史、协同分工", "敢于实验，也能回看每次变化"),
    ("README", "记录安装、运行、调参、标定与验收方法", "让代码可以被别人使用"),
    ("配置文件", "保存不同相机和光照下的参数", "把实验条件纳入版本管理"),
    ("日志/JSONL", "逐帧保存检测、时间和跟踪状态", "从“看画面”升级为“分析数据”"),
    ("模拟与离线录像", "没有底盘、MCU 和夹爪时独立开发视觉", "减少硬件等待，支持并行开发"),
    ("验收报告", "自动汇总检测比例、连续性、置信度和距离", "用统一标准讨论效果"),
], [1600, 4000, 3760])

add_heading(doc, "八、我们已经掌握什么，还欠缺什么", 1)
add_table(doc, ["层次", "目前具备", "仍需继续学习/验证"], [
    ("编程基础", "Python 模块、类、函数、条件、循环、文件读写、异常处理、命令行", "更系统的类型、测试设计、性能分析与重构"),
    ("视觉算法", "颜色分割、形态学、轮廓、形状特征、简单距离估计、时序平滑", "相机标定、畸变校正、姿态估计、复杂光照与学习方法"),
    ("工程开发", "模块化、参数化、离线运行、证据记录、验收指标", "持续集成、性能剖析、版本发布与长期维护"),
    ("机器人集成", "ROS2 节点、topic、消息、相机信息、统一输出接口", "真机时间同步、通信延迟、运动模糊、整车安全联调"),
    ("现场能力", "有验收清单和明确假设", "正式相机、安装位姿、场地光照、机械遮挡、真实速度下的测试"),
], [1300, 4030, 4030])

add_heading(doc, "九、适合低年级本科生的下一步路线", 1)
for text in [
    "能解释：用自己的话讲清楚每个脚本的输入、输出和用途。",
    "能复现：从新环境安装依赖，运行第 2 课，生成 scene、mask、annotated 和 result。",
    "能调参：选一段新录像，复制配置，调 HSV 与 ROI，并记录修改理由。",
    "能测量：在 0.3、0.5、0.8、1.0 m 采样，标定焦距并绘制误差表。",
    "能定位：故意设置错误参数，判断失败发生在配置、检测、时序还是输出层。",
    "能扩展：增加一种可配置颜色或一个新的形状筛选指标，并补充测试与文档。",
    "能集成：在 Ubuntu ROS2 中发布相机图像，检查 /cubes 字段、频率和时间戳。",
    "能验收：使用未参与调参的数据集生成报告，明确通过、失败和仍不确定的部分。",
]: add_number(doc, text)

add_heading(doc, "十、结语", 1)
doc.add_paragraph("对于只有少量编程基础的低年级本科生，这个视觉模块最大的价值，是让我们跨过了“会写几段代码”与“能参与真实项目”之间的第一道门槛。我们开始理解：算法需要接口，参数需要来源，结果需要证据，失败需要处理，未知需要记录，模块需要被别人运行，软件通过也不代表实机通过。")
add_callout(doc, "一句话总结", "我们学会的不是某几个 OpenCV API，而是一套可迁移的工程方法：把问题拆小，把逻辑分层，把参数外置，把结果记录，把正常和异常都测试，再用清晰接口接入完整系统。", PALE_GREEN)

# Source note
add_heading(doc, "附：本总结依据的项目材料", 2)
for text in [
    "docs/VISION_MODULE.md：视觉调参、距离标定、连续帧确认与验收标准。",
    "lessons/02_cube_perception：面向新手的独立视觉实验与输出样例。",
    "ros2_ws/src/cube_perception：检测器、时序过滤、独立工具、报告工具、ROS2 节点与配置。",
    "README.md 与 BEGINNER_PROJECT_LEARNING_GUIDE.md：系统模块边界、项目成熟度与初学者工程方法。",
]: add_bullet(doc, text)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
