from pathlib import Path
from docx import Document
from pypdf import PdfReader

downloads = Path(r"C:\Users\dahli\Downloads")
out = Path(__file__).resolve().parent / "_review_requirements_extract.txt"
parts = []

docx_path = downloads / "二审评分细则（竞技组）.docx"
doc = Document(docx_path)
parts.append("===== 二审评分细则（竞技组） =====")
for p in doc.paragraphs:
    if p.text.strip():
        parts.append(p.text.strip())
for ti, table in enumerate(doc.tables, 1):
    parts.append(f"--- 表格 {ti} ---")
    for row in table.rows:
        parts.append(" | ".join(cell.text.replace("\n", " / ").strip() for cell in row.cells))

pdf_path = downloads / "RoboGame2026 竞技组规则手册1_1.pdf"
reader = PdfReader(pdf_path)
parts.append("\n===== RoboGame2026 竞技组规则手册 =====")
for i, page in enumerate(reader.pages, 1):
    text = page.extract_text() or ""
    parts.append(f"\n--- 第 {i} 页 ---\n{text}")

text = "\n".join(parts)
try:
    repaired = text.encode("latin1").decode("utf-8")
    if repaired.count("比赛") + repaired.count("评分") > text.count("比赛") + text.count("评分"):
        text = repaired
except (UnicodeEncodeError, UnicodeDecodeError):
    pass
out.write_text(text, encoding="utf-8")
print(out)
